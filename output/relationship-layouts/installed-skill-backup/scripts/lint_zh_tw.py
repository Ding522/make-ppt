#!/usr/bin/env python3
"""Flag likely non-Taiwan wording without rewriting presentation content."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable


TERMS = {
    "雙擊": ("error", "按兩下／連按兩下"),
    "視頻": ("error", "影片"),
    "默認": ("error", "預設"),
    "用戶": ("error", "使用者"),
    "服務器": ("error", "伺服器"),
    "鼠標": ("error", "滑鼠"),
    "屏幕": ("error", "螢幕"),
    "鏈接": ("error", "連結"),
    "搜索": ("error", "搜尋"),
    "打印": ("error", "列印"),
    "內存": ("error", "記憶體"),
    "硬盤": ("error", "硬碟"),
    "加載": ("error", "載入"),
    "創建": ("error", "建立／新增"),
    "配置": ("review", "依語境使用設定、組態或配置"),
    "點擊": ("review", "按一下／點選"),
    "導入": ("review", "資料或檔案語境通常使用匯入"),
    "導出": ("review", "資料或檔案語境通常使用匯出"),
    "運行": ("review", "依語境使用執行或運作"),
    "支持": ("review", "能力或相容性語境通常使用支援"),
    "信息": ("review", "依語境使用資訊或訊息"),
    "數據": ("review", "一般資料語境通常使用資料"),
    "代碼": ("review", "程式內容通常使用程式碼；識別代碼可保留"),
    "組件": ("review", "軟體構成項目通常使用元件"),
    "保存": ("review", "儲存資料時使用儲存；保留之意可保留"),
}

VISIBLE_OUTLINE_SECTIONS = {
    "Deck Context Line",
    "Title",
    "Eyebrow",
    "Supporting Sentence",
    "Content",
    "Table Schema",
    "Takeaway",
}


@dataclass(frozen=True)
class Finding:
    path: str
    location: str
    severity: str
    term: str
    suggestion: str
    text: str


def strip_markdown_code(text: str) -> str:
    text = re.sub(r"`[^`]*`", "", text)
    text = re.sub(r"https?://\S+", "", text)
    return text


def markdown_lines(path: Path, visible_only: bool) -> Iterable[tuple[str, str]]:
    in_fence = False
    active_heading = ""
    for number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        stripped = raw_line.strip()
        if stripped.startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        heading = re.match(r"^(#{2,3})\s+(.+?)\s*$", stripped)
        if heading:
            active_heading = heading.group(2)
            if active_heading.startswith("Slide "):
                active_heading = ""
            continue
        if visible_only and active_heading not in VISIBLE_OUTLINE_SECTIONS:
            continue
        text = strip_markdown_code(raw_line).strip()
        if text:
            yield f"line {number}", text


def iter_shape_text(shape: object, prefix: str) -> Iterable[tuple[str, str]]:
    if getattr(shape, "has_text_frame", False):
        text = getattr(shape, "text", "").strip()
        if text:
            yield prefix, text
    if getattr(shape, "has_table", False):
        table = shape.table
        for row_index, row in enumerate(table.rows, 1):
            for column_index, cell in enumerate(row.cells, 1):
                text = cell.text.strip()
                if text:
                    yield f"{prefix} table R{row_index}C{column_index}", text
    for index, child in enumerate(getattr(shape, "shapes", ()), 1):
        yield from iter_shape_text(child, f"{prefix}.{index}")


def pptx_lines(path: Path) -> Iterable[tuple[str, str]]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    for slide_index, slide in enumerate(presentation.slides, 1):
        for shape_index, shape in enumerate(slide.shapes, 1):
            yield from iter_shape_text(shape, f"slide {slide_index} shape {shape_index}")


def load_allowlist(path: Path | None, inline: list[str]) -> set[str]:
    allowed = {term.strip() for term in inline if term.strip()}
    if path:
        for line in path.read_text(encoding="utf-8").splitlines():
            term = line.strip()
            if term and not term.startswith("#"):
                allowed.add(term)
    return allowed


def lint_path(path: Path, allowed: set[str], visible_only: bool) -> list[Finding]:
    if path.suffix.lower() in {".md", ".markdown"}:
        lines = markdown_lines(path, visible_only)
    elif path.suffix.lower() == ".pptx":
        lines = pptx_lines(path)
    else:
        raise ValueError(f"unsupported file type: {path}")

    findings = []
    for location, text in lines:
        for term, (severity, suggestion) in TERMS.items():
            if term in allowed or term not in text:
                continue
            findings.append(
                Finding(
                    path=str(path),
                    location=location,
                    severity=severity,
                    term=term,
                    suggestion=suggestion,
                    text=text.replace("\n", " / "),
                )
            )
    return findings


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", type=Path, help="Markdown outline or PPTX file")
    parser.add_argument(
        "--outline-visible-only",
        action="store_true",
        help="For Markdown, inspect only fields that become visible slide text.",
    )
    parser.add_argument("--allow-term", action="append", default=[], help="Allow one exact term")
    parser.add_argument("--allowlist", type=Path, help="UTF-8 file with one allowed term per line")
    parser.add_argument(
        "--fail-on-review",
        action="store_true",
        help="Return non-zero for contextual review findings as well as errors.",
    )
    parser.add_argument("--json", action="store_true", help="Emit machine-readable JSON")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    allowed = load_allowlist(args.allowlist, args.allow_term)
    findings: list[Finding] = []
    for path in args.paths:
        if not path.is_file():
            raise FileNotFoundError(path)
        findings.extend(lint_path(path, allowed, args.outline_visible_only))

    if args.json:
        print(json.dumps([asdict(item) for item in findings], ensure_ascii=False, indent=2))
    elif findings:
        for item in findings:
            print(
                f"{item.path}:{item.location}: {item.severity.upper()} "
                f"{item.term} → {item.suggestion} | {item.text}"
            )
    else:
        print("Taiwan terminology check: pass")

    blocking = any(item.severity == "error" for item in findings)
    if args.fail_on_review:
        blocking = blocking or bool(findings)
    return 1 if blocking else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(2)
