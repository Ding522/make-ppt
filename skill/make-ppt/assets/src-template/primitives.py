# primitives.py — reusable native-PowerPoint building blocks for the make-ppt grammar.
# Everything here produces EDITABLE PowerPoint objects (text boxes, autoshapes, connectors).
# Slide modules compose these; they never re-implement eyebrows, titles, takeaways,
# or semantic tables ad hoc.

import os
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
import theme as T

def _emu(v):
    """Coerce any length to an integer EMU.

    Centering math like `node_h / 2` turns an Emu (int subclass) into a Python
    float. python-pptx then serializes it as a non-integer EMU (e.g. `y="3360420.0"`,
    `cx="0.0"`). OOXML coordinates must be integers, so LibreOffice fails to parse
    them and collapses the shape/connector to the slide origin — this is why arrows
    and nodes "jump to the top of the slide". Every primitive that receives a
    geometry value funnels it through here so slide code can divide freely."""
    return Emu(int(round(float(v))))

def no_shadow(shape):
    """Strip any inherited drop shadow so every shape/line/picture renders flat.
    The make-ppt grammar is shadow-free by design — panels and cards use hairline
    borders, never shadows. Safe on autoshapes, connectors, and pictures. Any slide
    module that creates a shape directly (instead of via a primitive) must call this."""
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape

# ---------------------------------------------------------------- text core
def _set_run(r, text, *, color=T.INK, size=T.S_BODY, bold=False, mono=False, italic=False):
    r.text = text
    # snap_pt enforces the two type rules globally: min 10pt, and every size on the
    # 2pt grid — so even an off-grid literal in slide code is pulled into line.
    r.font.size = Pt(T.snap_pt(size))
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = T.FONT_MONO if mono else T.FONT_CJK
    # ensure the East Asian font is set too (python-pptx only sets latin by default)
    rPr = r._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', T.FONT_MONO if mono else T.FONT_CJK)

def rich_par(tf, segments, *, size=T.S_BODY, align=PP_ALIGN.LEFT, space_after=4, line=None, first=False):
    """Add one paragraph of mixed runs.
    segments: list of (text, opts) where opts may set color/bold/mono/size/italic.
    Mixed color runs in ONE paragraph = phrase-level title emphasis stays aligned."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    for text, opts in segments:
        o = {"color": T.INK, "size": size, "bold": False, "mono": False, "italic": False}
        o.update(opts or {})
        _set_run(p.add_run(), text, **o)
    return p

def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    x, y, w, h = _emu(x), _emu(y), _emu(w), _emu(h)
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

# ---------------------------------------------------------------- page scaffold
def add_background(slide, color=T.BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def add_eyebrow(slide, number, label, *, color=T.ORANGE, y=T.EYEBROW_Y):
    """Upper-left technical eyebrow:  // 16    MCP 推廣的阻力"""
    tb, tf = textbox(slide, T.MARGIN_X, y, Inches(7.5), Inches(0.3))
    rich_par(tf, [
        (f"// {number:02d}" if isinstance(number, int) else f"// {number}",
         {"color": color, "mono": True, "bold": True, "size": T.S_EYEBROW}),
        ("    " + label, {"color": color, "mono": True, "size": T.S_EYEBROW}),
    ], first=True)
    return tb

def _title_line_width_in(text, size_pt):
    """Rough one-line width of a heavy CJK title, in inches. CJK / full-width glyphs are
    ~1 em wide, latin / digits / spaces ~0.55 em — enough to decide one-line fit."""
    return sum((size_pt / 72.0) * (1.0 if ord(ch) > 0x2E80 else 0.55) for ch in text)

def fit_title_size(segments, box_w, *, size=T.S_TITLE, floor=T.S_TITLE_MIN):
    """Largest even size in [floor, size] at which the title fits on ONE line across
    box_w; steps down by 2pt. Keeps titles to a single line without hand-tuning."""
    text = "".join(t for t, _ in segments)
    box_in = float(box_w) / 914400.0
    s = int(size)
    while s > floor and _title_line_width_in(text, s) > box_in * 0.98:
        s -= T.STEP_PT
    return s

def add_argument_title(slide, segments, *, y=T.TITLE_Y, size=T.S_TITLE, w=None, line=1.08, fit=True):
    """Large editorial title. segments carry phrase-level emphasis:
       [("想從 Skill 對接", {}), ("全公司 MCP", {"color": T.ORANGE}), ("——卻卡在…", {})]
    fit=True keeps the title on ONE line: it starts at `size` (28pt) and steps down the
    2pt grid to S_TITLE_MIN only if the title is too long to fit one line — it never
    wraps unless even the floor size overflows. Write concise titles so 28pt holds."""
    box_w = w or T.CONTENT_W
    if fit:
        size = fit_title_size(segments, box_w, size=size)
    segs = [(t, dict({"bold": True, "size": size}, **(o or {}))) for t, o in segments]
    tb, tf = textbox(slide, T.MARGIN_X, y, box_w, Inches(1.2))
    rich_par(tf, segs, size=size, first=True, line=line)
    return tb

def add_supporting_sentence(slide, segments, *, y, w=None, size=T.S_SUPPORT):
    """Muted explanatory sentence under the title; bold key phrases stay INK."""
    segs = [(t, dict({"color": T.MUTED, "size": size}, **(o or {}))) for t, o in segments]
    tb, tf = textbox(slide, T.MARGIN_X, y, w or T.CONTENT_W, Inches(0.45))
    rich_par(tf, segs, size=size, first=True, line=1.25)
    return tb

def add_hairline(slide, y, *, x=None, w=None, color=T.HAIRLINE, weight=T.RULE_W):
    x = T.MARGIN_X if x is None else x
    w = T.CONTENT_W if w is None else w
    x, y, w = _emu(x), _emu(y), _emu(w)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
    ln.line.color.rgb = color
    ln.line.width = weight
    no_shadow(ln)
    return ln

def add_takeaway_line(slide, segments, *, note=None, y=None):
    """Explicit optional takeaway; never call it merely to fill the page bottom."""
    y = y if y is not None else T.SLIDE_H - Inches(0.82)
    add_hairline(slide, y, color=T.ORANGE, weight=T.ACCENT_RULE_W)
    segs = [(t, dict({"bold": True, "size": T.S_TAKEAWAY}, **(o or {}))) for t, o in segments]
    tb, tf = textbox(slide, T.MARGIN_X, y + Inches(0.12), T.CONTENT_W - Inches(2.2), Inches(0.5))
    rich_par(tf, segs, size=T.S_TAKEAWAY, first=True)
    if note:
        nb, nf = textbox(slide, T.SLIDE_W - T.MARGIN_X - Inches(2.6), y + Inches(0.16), Inches(2.6), Inches(0.35))
        rich_par(nf, [(note, {"color": T.MUTED, "size": 11, "mono": True})],
                 align=PP_ALIGN.RIGHT, first=True)


def _set_cell_border(cell, *, color=T.HAIRLINE, width=T.TABLE_BORDER_W):
    """Apply a flat hairline border to every edge of one native table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for index, edge_name in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        edge = tc_pr.find(qn(edge_name))
        if edge is None:
            edge = OxmlElement(edge_name)
            # CT_TableCellProperties puts line elements BEFORE cell fill.
            # PowerPoint ignores trailing line elements and falls back to the
            # table theme's white borders if this schema order is violated.
        else:
            tc_pr.remove(edge)
        tc_pr.insert(index, edge)
        edge.set("w", str(int(width)))
        for child in list(edge):
            edge.remove(child)
        solid_fill = OxmlElement("a:solidFill")
        color_node = OxmlElement("a:srgbClr")
        color_node.set("val", str(color))
        solid_fill.append(color_node)
        edge.append(solid_fill)
        edge.append(OxmlElement("a:prstDash"))
        edge[-1].set("val", "solid")


def _table_segments(value):
    if value is None:
        return [("", {})]
    if isinstance(value, str):
        return [(value, {})]
    if isinstance(value, (list, tuple)) and all(
        isinstance(item, (list, tuple)) and len(item) == 2 for item in value
    ):
        return [(str(text), dict(options or {})) for text, options in value]
    return [(str(value), {})]


def _table_alignment(value):
    if value in (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT):
        return value
    mapping = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }
    try:
        return mapping[str(value).lower()]
    except KeyError as exc:
        raise ValueError(f"Unsupported table alignment: {value}") from exc


def add_native_table(
    slide,
    x,
    y,
    w,
    h,
    *,
    headers,
    rows,
    col_widths=None,
    alignments=None,
    row_heights=None,
    row_label_col=None,
    merges=None,
):
    """Create one editable native PowerPoint table object.

    Use whenever content has column headers, repeated records, and cross-row
    alignment. Cell values may be strings or rich-run segment lists. `col_widths`
    are relative weights. `row_heights`, when supplied, contains absolute lengths
    for the header plus every body row. Merge coordinates are zero-based and include
    the header row: `(start_row, start_col, end_row, end_col)`.
    """
    headers = list(headers)
    rows = [list(row) for row in rows]
    if not headers:
        raise ValueError("Native tables require at least one header column")
    column_count = len(headers)
    if any(len(row) != column_count for row in rows):
        raise ValueError("Every native table row must match the header column count")

    x, y, w, h = _emu(x), _emu(y), _emu(w), _emu(h)
    row_count = len(rows) + 1
    frame = slide.shapes.add_table(row_count, column_count, x, y, w, h)
    table = frame.table
    no_shadow(frame)

    widths = list(col_widths or [1] * column_count)
    if len(widths) != column_count or any(float(value) <= 0 for value in widths):
        raise ValueError("col_widths must contain one positive weight per column")
    total_weight = sum(float(value) for value in widths)
    assigned = 0
    for index, weight in enumerate(widths):
        width = w - assigned if index == column_count - 1 else _emu(w * float(weight) / total_weight)
        table.columns[index].width = width
        assigned += width

    if row_heights is not None:
        heights = [_emu(value) for value in row_heights]
        if len(heights) != row_count:
            raise ValueError("row_heights must contain header height plus every body row")
    else:
        header_height = min(_emu(T.TABLE_HEADER_H), h)
        body_height = _emu((h - header_height) / max(1, len(rows)))
        heights = [header_height] + [body_height] * len(rows)
        if len(rows):
            heights[-1] += h - sum(heights)
    for index, height in enumerate(heights):
        table.rows[index].height = height

    for start_row, start_col, end_row, end_col in merges or []:
        table.cell(start_row, start_col).merge(table.cell(end_row, end_col))

    column_alignments = list(alignments or ["left"] * column_count)
    if len(column_alignments) != column_count:
        raise ValueError("alignments must contain one value per column")

    values = [headers] + rows
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            if getattr(cell, "is_spanned", False):
                continue

            is_header = row_index == 0
            is_row_label = (
                not is_header
                and row_label_col is not None
                and column_index == row_label_col
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = T.INK if is_header else T.BG
            cell.margin_left = cell.margin_right = T.TABLE_CELL_MARGIN_X
            cell.margin_top = cell.margin_bottom = T.TABLE_CELL_MARGIN_Y
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _set_cell_border(cell)

            base_options = {
                "color": T.ON_DARK if is_header else T.INK_SOFT,
                "size": T.S_TABLE_HEADER if is_header else T.S_TABLE_BODY,
                "bold": is_header or is_row_label,
            }
            segments = [
                (text, dict(base_options, **options))
                for text, options in _table_segments(value)
            ]
            text_frame = cell.text_frame
            text_frame.clear()
            text_frame.word_wrap = True
            rich_par(
                text_frame,
                segments,
                size=base_options["size"],
                align=_table_alignment(column_alignments[column_index]),
                first=True,
                space_after=0,
                line=1.05,
            )

    return frame

# ---------------------------------------------------------------- shapes / motifs
def add_panel(slide, x, y, w, h, *, fill=T.PANEL_BG, line=T.PANEL_LN, radius=True, line_w=Pt(1)):
    x, y, w, h = _emu(x), _emu(y), _emu(w), _emu(h)
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        try:
            shp.adjustments[0] = T.RADIUS_SMALL
        except Exception:
            pass
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = line_w
    no_shadow(shp)
    return shp

def add_status_pill(slide, x, y, text, *, bg=T.GREEN, fg=None, w=Inches(0.62), h=Inches(0.26), size=T.S_LABEL, mono=False):
    """Small semantic label: 期待 / 現實 / RISK / EPIC …"""
    shp = add_panel(slide, x, y, w, h, fill=bg, line=None)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    rich_par(tf, [(text, {"color": fg or T.ON_DARK, "bold": True, "size": size, "mono": mono})],
             align=PP_ALIGN.CENTER, first=True)
    return shp

def add_chip(slide, x, y, text, *, color=T.INK, fill=T.PANEL_BG, h=Inches(0.28),
             size=9.5, mono=True, line_w=Pt(1)):
    """Single-token component label (ADK / Sandbox / Gateway / Model Armor…).

    Sized to its own text and — crucially — word-wrap is OFF, so a short label can
    never break mid-word into "AD / K". Width counts CJK glyphs as double and adds the
    inner margins, so caps-heavy tokens still fit. Returns the chip width so callers
    can lay chips out in a row. Prefer this over hand-rolling a rounded rectangle."""
    x, y, h = _emu(x), _emu(y), _emu(h)
    inner = Inches(0.07)
    units = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
    text_w = Emu(int(units * (0.62 * size) / 72 * 914400))  # pt advance → inch → EMU
    w = _emu(text_w + inner * 2 + Inches(0.10))
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_w is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = color; shp.line.width = line_w
    no_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = False
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    tf.margin_left = tf.margin_right = inner
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rich_par(tf, [(text, {"color": color, "mono": mono, "size": size})],
             align=PP_ALIGN.CENTER, first=True)
    return w

def add_label_panel(slide, x, y, w, h, segments, *, fill=T.PANEL_BG, line=T.HAIRLINE,
                    align=PP_ALIGN.CENTER, radius=True, line_w=Pt(1)):
    """Bordered panel with its label centered INSIDE it (vertical + horizontal).

    Use this instead of `add_panel(...)` followed by a separate `textbox(...)`: that
    hand-rolled pattern leaves the label anchored to the TOP of the panel (the default),
    so text like a "90% / 10%" badge floats at the top edge instead of centering."""
    shp = add_panel(slide, x, y, w, h, fill=fill, line=line, radius=radius, line_w=line_w)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    rich_par(tf, [(t, dict({}, **(o or {}))) for t, o in segments], align=align, first=True)
    return shp

def add_step_marker(slide, x, y, n, *, d=Inches(0.3), bg=T.ORANGE, fg=None):
    """Numbered circle for sequences and workflows."""
    x, y, d = _emu(x), _emu(y), _emu(d)
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    c.fill.solid(); c.fill.fore_color.rgb = bg; c.line.fill.background()
    no_shadow(c)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    rich_par(tf, [(str(n), {"color": fg or T.ON_DARK, "bold": True, "size": 11, "mono": True})],
             align=PP_ALIGN.CENTER, first=True)
    return c

def add_arrow(slide, x1, y1, x2, y2, *, color=T.INK_SOFT, weight=Pt(1.5), dashed=False, arrow=True):
    # Coerce to integer EMU: a connector fed a float coordinate (from `.../ 2`
    # centering) serializes as a non-integer EMU that LibreOffice cannot parse,
    # which throws the line to the slide origin. See _emu().
    x1, y1, x2, y2 = _emu(x1), _emu(y1), _emu(x2), _emu(y2)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.line.color.rgb = color
    ln.line.width = weight
    if dashed:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        # arrowhead at the tail end (dashed background/feedback lines often omit it)
        lnEl = ln.line._get_or_add_ln()
        tail = lnEl.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        lnEl.append(tail)
    no_shadow(ln)
    return ln

def add_technical_node(slide, x, y, w, h, title, *, sub=None, fill=T.PANEL_BG,
                       line=T.INK, title_color=T.INK, mono_title=False, line_w=Pt(1.25)):
    """Bordered system node for architecture diagrams."""
    shp = add_panel(slide, x, y, w, h, fill=fill, line=line, radius=False, line_w=line_w)
    tf = shp.text_frame
    tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.08)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rich_par(tf, [(title, {"bold": True, "size": 13, "color": title_color, "mono": mono_title})], first=True, space_after=1)
    if sub:
        rich_par(tf, [(sub, {"color": T.MUTED, "size": 10, "mono": True})], space_after=0)
    return shp

def add_terminal_panel(slide, x, y, w, h, title, lines, *, badge=None, badge_bg=T.RED):
    """Dark terminal/config evidence panel with mac traffic lights.
    lines: list of paragraphs; each paragraph = list of (text, opts) mono runs.
    Highlight secrets with {"color": T.RED} or chip-like {"color": T.ORANGE}."""
    panel = add_panel(slide, x, y, w, h, fill=T.DARK_PANEL, line=None)
    # traffic lights
    for i, c in enumerate((T.RED, T.ORANGE, T.GREEN)):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                     x + Inches(0.16 + i * 0.22), y + Inches(0.14),
                                     Inches(0.11), Inches(0.11))
        dot.fill.solid(); dot.fill.fore_color.rgb = c; dot.line.fill.background()
        no_shadow(dot)
    tb, tf = textbox(slide, x + Inches(0.85), y + Inches(0.1), w - Inches(1.9), Inches(0.25))
    rich_par(tf, [(title, {"color": T.DARK_PANEL_MUTED, "mono": True, "size": 10.5})], first=True)
    if badge:
        add_status_pill(slide, x + w - Inches(0.75), y + Inches(0.11), badge,
                        bg=badge_bg, w=Inches(0.56), h=Inches(0.22), size=9, mono=True)
    bb, bf = textbox(slide, x + Inches(0.22), y + Inches(0.48), w - Inches(0.44), h - Inches(0.6))
    for i, line in enumerate(lines):
        segs = [(t, dict({"color": T.DARK_PANEL_TXT, "mono": True, "size": T.S_MONO}, **(o or {})))
                for t, o in line]
        rich_par(bf, segs, size=T.S_MONO, first=(i == 0), space_after=3, line=1.15)
    return panel

def add_metric_callout(slide, x, y, segments, *, size=54, w=Inches(8), caption=None):
    """Huge KPI / tension figure, e.g.  66% → 97%   or   5x → 2xx → 1xxx"""
    x, y = _emu(x), _emu(y)
    tb, tf = textbox(slide, x, y, w, Inches(1.2))
    segs = [(t, dict({"bold": True, "size": size, "mono": True}, **(o or {}))) for t, o in segments]
    rich_par(tf, segs, size=size, first=True)
    if caption:
        cb, cf = textbox(slide, x, y + Inches(0.16 + size / 72), w, Inches(0.35))
        rich_par(cf, [(caption, {"color": T.MUTED, "size": 12})], first=True)
    return tb

# ---------------------------------------------------------------- full-slide scaffolds
def add_cover_slide(slide, title_segments, *, variant="editorial-light", context=None,
                    eyebrow=None, subtitle=None, presenter=None, affiliation=None,
                    date=None, page_marker=None, footer_highlight=None):
    """Build one of the two supported cover compositions.

    ``editorial-light`` follows the primary reference cover: warm off-white canvas,
    oversized two-line editorial title, short orange rule, optional subtitle, and a
    quiet author/date footer. ``report-dark`` is the formal monthly/progress/
    performance-report cover: near-black canvas, white title, orange metadata, and an
    optional bottom scope/highlight line.

    ``footer_highlight`` is truly optional. Never invent a fixed count or summary such
    as "三件事" merely to fill the bottom of a dark cover.
    """
    if variant not in {"editorial-light", "report-dark"}:
        raise ValueError(f"unknown cover variant: {variant}")

    context = context or T.DECK_CONTEXT
    is_dark = variant == "report-dark"
    background = T.INK if is_dark else T.BG
    primary = T.ON_DARK if is_dark else T.INK
    secondary = T.DARK_PANEL_MUTED if is_dark else T.MUTED
    chrome = T.ON_DARK if is_dark else T.INK

    add_background(slide, background)
    add_hairline(slide, Inches(0.5), color=chrome, weight=Pt(1.25))

    cb, cf = textbox(slide, T.MARGIN_X, Inches(0.66), Inches(7.4), Inches(0.3))
    rich_par(cf, [(context, {"color": chrome, "mono": True, "size": T.S_MONO})], first=True)

    header_right = page_marker or (date if is_dark else None)
    if header_right:
        hb, hf = textbox(
            slide,
            T.SLIDE_W - T.MARGIN_X - Inches(2.5),
            Inches(0.66),
            Inches(2.5),
            Inches(0.3),
        )
        rich_par(
            hf,
            [(header_right, {"color": secondary, "mono": True, "size": T.S_MONO})],
            align=PP_ALIGN.RIGHT,
            first=True,
        )

    if is_dark:
        if eyebrow:
            eb, ef = textbox(slide, T.MARGIN_X, Inches(1.55), Inches(8.4), Inches(0.34))
            rich_par(
                ef,
                [(eyebrow, {"color": T.ORANGE, "mono": True, "bold": True, "size": T.S_EYEBROW})],
                first=True,
            )

        title_y = Inches(2.12) if eyebrow else Inches(1.9)
        title_opts = [
            (text, dict({"color": primary, "bold": True, "size": T.S_COVER_TITLE_DARK}, **(opts or {})))
            for text, opts in title_segments
        ]
        tb, tf = textbox(slide, T.MARGIN_X, title_y, Inches(11.2), Inches(1.6), wrap=True)
        rich_par(tf, title_opts, size=T.S_COVER_TITLE_DARK, first=True, line=1.02)

        accent_y = Inches(3.82)
        add_hairline(slide, accent_y, w=Inches(1.55), color=T.ORANGE, weight=Pt(3.0))
        if subtitle:
            sb, sf = textbox(slide, T.MARGIN_X, Inches(4.05), Inches(10.7), Inches(0.6))
            rich_par(
                sf,
                [(subtitle, {"color": secondary, "size": T.S_COVER_SUBTITLE})],
                first=True,
                line=1.2,
            )

        if presenter or affiliation:
            ab, af = textbox(slide, T.MARGIN_X, Inches(4.8), Inches(10.8), Inches(0.42))
            author_segments = []
            if presenter:
                author_segments.append((presenter, {"color": primary, "bold": True, "size": T.S_SUPPORT}))
            if presenter and affiliation:
                author_segments.append(("   |   ", {"color": secondary, "size": T.S_SUPPORT}))
            if affiliation:
                author_segments.append((affiliation, {"color": secondary, "size": T.S_SUPPORT}))
            rich_par(af, author_segments, first=True)

        add_hairline(slide, Inches(5.35), color=T.DARK_PANEL_MUTED)
        if footer_highlight:
            fb, ff = textbox(slide, T.MARGIN_X, T.SLIDE_H - Inches(0.88), Inches(11.5), Inches(0.38))
            rich_par(
                ff,
                [(footer_highlight, {"color": T.ORANGE, "mono": True, "bold": True, "size": T.S_SUPPORT})],
                first=True,
            )
    else:
        if eyebrow:
            eb, ef = textbox(slide, T.MARGIN_X, Inches(1.98), Inches(8.4), Inches(0.34))
            rich_par(
                ef,
                [(eyebrow, {"color": T.ORANGE, "mono": True, "size": T.S_EYEBROW})],
                first=True,
            )

        title_y = Inches(2.58) if eyebrow else Inches(2.32)
        title_opts = [
            (text, dict({"color": primary, "bold": True, "size": T.S_COVER_TITLE_LIGHT}, **(opts or {})))
            for text, opts in title_segments
        ]
        tb, tf = textbox(slide, T.MARGIN_X, title_y, Inches(9.4), Inches(2.0), wrap=True)
        rich_par(tf, title_opts, size=T.S_COVER_TITLE_LIGHT, first=True, line=0.94)

        add_hairline(slide, Inches(4.92), w=Inches(1.55), color=T.ORANGE, weight=Pt(3.0))
        if subtitle:
            sb, sf = textbox(slide, T.MARGIN_X, Inches(5.28), Inches(9.4), Inches(0.92))
            rich_par(
                sf,
                [(subtitle, {"color": secondary, "size": T.S_COVER_SUBTITLE})],
                first=True,
                line=1.25,
            )

        add_hairline(slide, Inches(6.63), color=T.HAIRLINE)
        if presenter:
            pb, pf = textbox(slide, T.MARGIN_X, Inches(6.83), Inches(6.0), Inches(0.3))
            rich_par(
                pf,
                [(presenter, {"color": primary, "mono": True, "size": T.S_MONO})],
                first=True,
            )
        if date:
            db, df = textbox(
                slide,
                T.SLIDE_W - T.MARGIN_X - Inches(2.5),
                Inches(6.83),
                Inches(2.5),
                Inches(0.3),
            )
            rich_par(
                df,
                [(date, {"color": secondary, "mono": True, "size": T.S_MONO})],
                align=PP_ALIGN.RIGHT,
                first=True,
            )


def add_section_slide(slide, number, part_label, title_segments, bullets, nav, active_idx,
                      *, corner_badge=None, context=None):
    """Dark section divider: giant orange number, PART label, white title with orange
    emphasis, bullet preview, bottom nav of all parts with the active one lit.

    `context` is the mono line at the top-left: the deck's REAL context (topic / team /
    date), defaulting to T.DECK_CONTEXT. Never pass invented event or summit branding —
    this deck is shared with colleagues, not hosted at a conference."""
    add_background(slide, T.INK)
    add_hairline(slide, Inches(0.5), color=T.ON_DARK, weight=Pt(1.25))
    tb, tf = textbox(slide, T.MARGIN_X, Inches(0.66), Inches(6), Inches(0.3))
    rich_par(tf, [(context or T.DECK_CONTEXT,
                   {"color": T.ON_DARK, "mono": True, "size": T.S_MONO})], first=True)
    hb, hf = textbox(slide, T.SLIDE_W - T.MARGIN_X - Inches(2.5), Inches(0.66), Inches(2.5), Inches(0.3))
    rich_par(hf, [(corner_badge or "SECTION", {"color": T.DARK_PANEL_MUTED, "mono": True, "size": 12})],
             align=PP_ALIGN.RIGHT, first=True)
    nb, nf = textbox(slide, T.MARGIN_X, Inches(2.35), Inches(3.9), Inches(2.6), wrap=False)
    rich_par(nf, [(f"{number:02d}", {"color": T.ORANGE, "bold": True, "size": T.S_SECTION_NUM, "mono": True})], first=True)
    pb, pf = textbox(slide, Inches(4.1), Inches(2.35), Inches(8.4), Inches(0.35))
    rich_par(pf, [(part_label, {"color": T.ORANGE, "mono": True, "bold": True, "size": 15})], first=True)
    ttb, ttf = textbox(slide, Inches(4.1), Inches(2.75), Inches(8.6), Inches(1.0))
    segs = [(t, dict({"bold": True, "size": 38, "color": T.ON_DARK}, **(o or {}))) for t, o in title_segments]
    rich_par(ttf, segs, first=True)
    bb, bf = textbox(slide, Inches(4.3), Inches(3.85), Inches(8.2), Inches(2.2))
    for i, b in enumerate(bullets):
        rich_par(bf, [("·  ", {"color": T.ORANGE, "size": 13, "bold": True}),
                      (b, {"color": RGB_soft_on_dark(), "size": 13})],
                 first=(i == 0), space_after=7)
    xs = [T.MARGIN_X, Inches(4.9), Inches(9.0)]
    for i, label in enumerate(nav):
        on = (i == active_idx)
        vb, vf = textbox(slide, xs[i], T.SLIDE_H - Inches(0.66), Inches(3.8), Inches(0.3))
        rich_par(vf, [(f"{i+1:02d}  ", {"color": T.ORANGE if on else T.DARK_PANEL_MUTED, "mono": True, "bold": on, "size": 12}),
                      (label, {"color": T.ON_DARK if on else T.DARK_PANEL_MUTED, "size": 12, "bold": on})], first=True)

def RGB_soft_on_dark():
    from pptx.dml.color import RGBColor
    return RGBColor(0xCF, 0xC9, 0xC0)

def add_screenshot(slide, img_path, x, y, *, max_w, max_h, frame=True):
    """Place a raster source screenshot preserving aspect ratio (fit inside max box).

    If the file is missing, reserve the space with a placeholder instead of crashing
    the whole build (a single absent asset must not take the deck down)."""
    x, y, max_w, max_h = _emu(x), _emu(y), _emu(max_w), _emu(max_h)
    if not img_path or not os.path.exists(str(img_path)):
        return add_photo_slot(slide, x, y, max_w, max_h, label="圖片缺檔", frame=frame)
    from PIL import Image
    iw, ih = Image.open(img_path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic = slide.shapes.add_picture(img_path, x, y, width=w, height=h)
    if frame:
        pic.line.color.rgb = T.PANEL_LN
        pic.line.width = Pt(1)
    no_shadow(pic)
    return pic

def add_photo_slot(slide, x, y, w, h, *, img_path=None, caption=None, label="圖片待補", frame=True):
    """Reserve a fixed region for a photo the user fills in later.

    The default (no img_path) draws a flat, dashed hairline placeholder with a muted
    mono label — the layout keeps the exact space, and the user drops an image into
    that box in PowerPoint afterwards (or sets img_path and rebuilds). If img_path is
    given AND exists, the photo is placed fit-inside, aspect preserved.

    This is the standard way to add on-site / illustrative photos in this skill: the
    deck reserves slots and the user decides what (if anything) goes in them. Genuine
    evidence screenshots that PROVE a claim (Kanban, logs, metrics) still go through
    add_screenshot with a real path."""
    x, y, w, h = _emu(x), _emu(y), _emu(w), _emu(h)
    if img_path and os.path.exists(str(img_path)):
        placed = add_screenshot(slide, str(img_path), x, y, max_w=w, max_h=h, frame=frame)
    else:
        box = add_panel(slide, x, y, w, h, fill=T.BG, line=T.HAIRLINE, radius=False)
        try:
            box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        except Exception:
            pass
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        rich_par(tf, [(label, {"color": T.MUTED, "mono": True, "size": 9})],
                 align=PP_ALIGN.CENTER, first=True)
        placed = box
    if caption:
        cb, cf = textbox(slide, x, y + h + Inches(0.03), w, Inches(0.2))
        rich_par(cf, [(caption, {"color": T.MUTED, "mono": True, "size": 8})],
                 align=PP_ALIGN.CENTER, first=True)
    return placed
