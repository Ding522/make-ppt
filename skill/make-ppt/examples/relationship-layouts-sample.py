"""Render optional layouts using excerpts from the user's visual references.

Run: python relationship-layouts-sample.py OUTPUT.pptx
These reference excerpts demonstrate layout only; never reuse as new-deck evidence.
"""
import sys
from pathlib import Path
from datetime import date
from pptx import Presentation
from pptx.util import Inches as I

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'assets' / 'src-template'))
import theme as T
import primitives as P
import relationship_layouts as R


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = T.SLIDE_W, T.SLIDE_H
    def slide(title, support):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        P.add_background(s)
        P.add_eyebrow(s, len(prs.slides), '版型示範 · 參考圖內容節錄')
        parts = title.split('：', 1) if '：' in title else title.split('，', 1)
        P.add_argument_title(s, [(parts[0], {'bold': True})] +
            ([('：' + parts[1], {'bold': True, 'color': T.ORANGE})] if len(parts)>1 else []))
        P.add_supporting_sentence(s, [(support, {})], y=I(1.47))
        P.add_hairline(s, I(1.95))
        return s
    s = slide('到職九個月，三條主線接力推進', '共用時間軸呈現工作期間；分開的色帶保留實際間隔。')
    d = lambda year, month: date(year, month, 1).toordinal()
    R.add_multitrack_timeline(s, T.MARGIN_X, I(2.2), T.CONTENT_W, I(3.8),
        ticks=[(d(2025,10),'2025/10'),(d(2025,11),'11'),(d(2025,12),'12'),
               (d(2026,1),'2026/01'),(d(2026,2),'02'),(d(2026,3),'03'),
               (d(2026,4),'04'),(d(2026,5),'05'),(d(2026,6),'06'),(d(2026,7),'07')],
        tracks=[
            dict(label='暖心陪伴系統 App', color=T.BLUE,
                 intervals=[(d(2025,10),d(2026,2))], note='Reminder 後端 / Cloud SQL · SOS · FCM 解耦重構'),
            dict(label='AI 架構師助手', color=T.ORANGE,
                 intervals=[(d(2026,2),d(2026,3)),(d(2026,6),d(2026,7))],
                 note='model 選型 / 產圖 pipeline · Data Lake DNS 修復 · Fortify 190 → 26'),
            dict(label='雲端處 ITSM 平台', color=T.GREEN,
                 intervals=[(d(2026,4),d(2026,7))], note='資訊系統 / 機房 schema · BCM 戰情室 / VM 納管')])
    s = slide('十項能力 × 五條專案線', '左側色條：藍＝功能層、橘＝系統層、綠＝複用層；格內符號見下方圖例。')
    names=['S1 行動端 App 開發（Flutter）','S2 後端服務與全端交付','S3 正確性與效能工程',
           'S4 資料建模與資料庫架構','S5 資料治理與 ETL','S6 雲端與微服務架構取捨',
           'S7 資安合規落地','S8 流程設計與平台化','S9 AI 應用開發','S10 用 AI 改工作流程']
    patterns=['DDDuu','DDDOD','DDDDD','uDDOD','uuuuD','DODDO','DuuDD','uuuOD','uuuDu','uuuDD']
    states={'D':dict(symbol='●',label='深度投入',color=T.ORANGE),
            'O':dict(symbol='○',label='有參與',color=T.MUTED),
            'u':dict(symbol='·',label='參考圖未標示',color=T.HAIRLINE)}
    R.add_status_matrix(s,T.MARGIN_X,I(2.12),T.CONTENT_W,I(4.75),
        headers=['能力','平安守護鈴','日日好生活','FCM 推播','AI 架構師助手','ITSM 平台'],
        rows=[[name]+list(pattern) for name,pattern in zip(names,patterns)],states=states,
        col_widths=[3.3,1.55,1.3,1.3,1.55,1.3], groups=[
            dict(start=0,end=1,color=T.BLUE),dict(start=2,end=6,color=T.ORANGE),
            dict(start=7,end=9,color=T.GREEN)])
    s = slide('十項能力分屬三層：功能層 / 系統層 / 複用層', '各層以色條分組，能力名稱與主要證據左右對照。')
    evidence=['日日好生活、平安守護鈴','sectest · BCM · reminder / notification',
              'reminder · FCM · 設備清單 · 架構師助手','15 schema / 78 表 · BCM 4 表',
              'DeviceInfo · ICTINV · LDAP','SOS 三層前濾 · 推播解耦 · Data Lake',
              '模型選型 · bcrypt / EasyAuth · RBAC','BCM 檢核流程 · 黑／灰箱複用',
              'AI 架構師助手','SDD · Fortify 迴圈 · 文件流水線']
    R.add_layered_list(s,T.MARGIN_X,I(2.12),T.CONTENT_W,I(4.65),groups=[
        dict(label='層一 · 功能層',description='把一條功能完整做出來',color=T.BLUE,items=list(zip(names[:2],evidence[:2]))),
        dict(label='層二 · 系統層',description='讓系統正確、可維護、安全',color=T.ORANGE,items=list(zip(names[2:7],evidence[2:7]))),
        dict(label='層三 · 複用層',description='讓做法能被重複使用',color=T.GREEN,items=list(zip(names[7:],evidence[7:])))])
    # A plain table shows that prose stays prose; no forced conversion into dots.
    s = slide('一般文字表格仍保留完整說明', '這頁比較版型用途，使用原有標準表格即可。')
    P.add_native_table(s,T.MARGIN_X,I(2.3),T.CONTENT_W,I(2.6),
        headers=['呈現方式','適用內容'],rows=[['多軌時間帶','專案期間、重疊與接續關係'],
        ['狀態矩陣','兩個分類軸之間的離散狀態'],['分層清單','分類下的項目與對應證據']],
        col_widths=[1,3])
    output=Path(output); output.parent.mkdir(parents=True,exist_ok=True)
    prs.save(output)


if __name__ == '__main__':
    build(sys.argv[1])
