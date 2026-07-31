#!/usr/bin/env python3
"""Build a source-grounded, incrementally cached evidence pack.

Usage:
  build_evidence_pack.py <source_root> <output_dir> [--force]

The pack stores one extracted Markdown file per supported source plus a stable
manifest. Unchanged files reuse their extracted artifact by SHA-256, so repeated
presentation runs do not re-parse the same source.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import zipfile
from pathlib import Path
from typing import Any


SUPPORTED = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".md": "markdown",
    ".txt": "text",
    ".csv": "csv",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".xlsx": "xlsx",
}
IGNORED_DIRS = {
    ".git",
    ".scratch",
    ".temp",
    "__pycache__",
    "dist",
    "node_modules",
    "output",
    "preview",
    "rendered",
    "presentation",
    "tmp",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_name(relative: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", relative.replace("\\", "/"))
    return name.strip("._") or "source"


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks = []
    for number, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            value = getattr(shape, "text", "").strip()
            if value:
                lines.append(value)
        blocks.append(f"## Slide {number}\n\n" + ("\n\n".join(lines) or "(no text)"))
    return "\n\n".join(blocks) or "(empty presentation)"


def extract_docx(path: Path) -> str:
    try:
        from docx import Document

        doc = Document(str(path))
        parts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        for index, table in enumerate(doc.tables, 1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            parts.append(f"## Table {index}\n\n" + "\n".join(rows))
        return "\n\n".join(parts) or "(empty document)"
    except ImportError:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        return re.sub(r"<[^>]+>", " ", xml).replace("&amp;", "&").strip() or "(empty document)"


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=False, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(" | ".join(values))
        parts.append(f"## Sheet: {sheet.title}\n\n" + ("\n".join(rows) or "(empty sheet)"))
    return "\n\n".join(parts) or "(empty workbook)"


def extract_pdf(path: Path) -> tuple[str, str]:
    try:
        result = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            check=False,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.replace("\r\n", "\n"), "pdftotext"
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    for module_name in ("pypdf", "PyPDF2"):
        try:
            module = __import__(module_name)
            reader = module.PdfReader(str(path))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
            if text:
                return text, module_name
        except (ImportError, OSError, ValueError):
            continue
    return "(PDF text extraction unavailable; inspect the source pages directly.)", "unavailable"


def extract_source(path: Path, kind: str) -> tuple[str, str, dict[str, Any]]:
    if kind in {"markdown", "text", "csv"}:
        return read_text(path), "plain-text", {}
    if kind == "pptx":
        return extract_pptx(path), "python-pptx", {}
    if kind == "docx":
        return extract_docx(path), "python-docx or OOXML fallback", {}
    if kind == "xlsx":
        return extract_xlsx(path), "openpyxl", {}
    if kind == "pdf":
        text, extractor = extract_pdf(path)
        return text, extractor, {}
    if kind == "image":
        try:
            from PIL import Image

            with Image.open(path) as image:
                width, height = image.size
            return (
                "(Raster asset; inspect the image when visual evidence is relevant.)",
                "Pillow metadata",
                {"width": width, "height": height},
            )
        except (ImportError, OSError):
            return "(Raster asset; inspect the image when visual evidence is relevant.)", "metadata unavailable", {}
    raise ValueError(f"unsupported source kind: {kind}")


def discover(root: Path, excluded_paths: tuple[Path, ...] = ()) -> list[Path]:
    files = []
    for path in root.rglob("*"):
        if not path.is_file() or path.name.startswith("~$") or path.name == ".DS_Store":
            continue
        resolved = path.resolve()
        if any(resolved == excluded or excluded in resolved.parents for excluded in excluded_paths):
            continue
        if any(part in IGNORED_DIRS for part in path.relative_to(root).parts[:-1]):
            continue
        if path.suffix.lower() in SUPPORTED:
            files.append(path)
    return sorted(files, key=lambda item: item.relative_to(root).as_posix().lower())


def render_source_markdown(relative: str, digest: str, kind: str, extractor: str,
                          metadata: dict[str, Any], content: str) -> str:
    metadata_lines = "".join(f"- {key}: {value}\n" for key, value in metadata.items())
    return (
        f"# Source: {relative}\n\n"
        f"- Relative path: `{relative}`\n"
        f"- SHA-256: `{digest}`\n"
        f"- Kind: `{kind}`\n"
        f"- Extractor: `{extractor}`\n"
        f"{metadata_lines}\n"
        "## Extracted Evidence\n\n"
        f"{content.rstrip()}\n"
    )


def build_pack(root: Path, output: Path, force: bool) -> tuple[int, int, int]:
    output.mkdir(parents=True, exist_ok=True)
    sources_dir = output / "sources"
    sources_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output / "manifest.json"
    old: dict[str, Any] = {}
    if manifest_path.exists() and not force:
        try:
            old = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            old = {}

    entries = []
    reused = extracted = 0
    active_files = set()
    for path in discover(root, (output,)):
        relative = path.relative_to(root).as_posix()
        kind = SUPPORTED[path.suffix.lower()]
        digest = sha256(path)
        previous = next((item for item in old.get("sources", []) if item.get("path") == relative), None)
        content_file = previous.get("content_file") if previous and previous.get("sha256") == digest else None
        if content_file and (output / content_file).exists():
            reused += 1
            extractor = previous.get("extractor", "cached")
            metadata = previous.get("metadata", {})
        else:
            content, extractor, metadata = extract_source(path, kind)
            content_file = f"sources/{digest[:12]}-{safe_name(relative)}.md"
            (output / content_file).write_text(
                render_source_markdown(relative, digest, kind, extractor, metadata, content),
                encoding="utf-8",
                newline="\n",
            )
            extracted += 1
        active_files.add(content_file)
        entries.append({
            "path": relative,
            "kind": kind,
            "sha256": digest,
            "content_file": content_file,
            "extractor": extractor,
            "metadata": metadata,
        })

    removed = 0
    for candidate in sources_dir.glob("*.md"):
        relative = candidate.relative_to(output).as_posix()
        if relative not in active_files:
            candidate.unlink()
            removed += 1

    manifest = {"schema_version": 1, "source_root": ".", "sources": entries}
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [
        "# Source Evidence Pack",
        "",
        "This pack is the planner's complete, source-grounded reading set.",
        "Unchanged sources reuse extracted Markdown by SHA-256.",
        "",
        "## Sources",
        "",
    ]
    for entry in entries:
        lines.append(
            f"- `{entry['path']}` ({entry['kind']}) → `" 
            f"{entry['content_file']}` · `{entry['sha256'][:12]}`"
        )
    if not entries:
        lines.append("- No supported sources found.")
    (output / "manifest.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    return len(entries), reused, removed


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_root", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--force", action="store_true", help="re-extract unchanged sources")
    args = parser.parse_args()
    root = args.source_root.resolve()
    output = args.output_dir.resolve()
    if not root.is_dir():
        parser.error(f"source root is not a directory: {root}")
    total, reused, removed = build_pack(root, output, args.force)
    extracted = total - reused
    print(f"evidence pack: {output} ({total} sources, {reused} reused, {extracted} extracted, {removed} stale removed)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
