#!/usr/bin/env python3
"""Validate object-model requirements that image previews cannot prove."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation


def parse_slide_spec(spec: str, slide_count: int) -> list[int]:
    selected: set[int] = set()
    for token in filter(None, (part.strip() for part in spec.split(","))):
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            start, end = int(start_text), int(end_text)
            if start > end:
                raise ValueError(f"Invalid slide range: {token}")
            selected.update(range(start, end + 1))
        else:
            selected.add(int(token))
    invalid = sorted(number for number in selected if number < 1 or number > slide_count)
    if invalid:
        raise ValueError(
            f"Slide selection out of range for {slide_count}-slide deck: {invalid}"
        )
    return sorted(selected)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("deck", type=Path)
    parser.add_argument(
        "--require-table-slides",
        default="",
        help='One-based slide numbers that must contain native tables, e.g. "3,7-9".',
    )
    args = parser.parse_args()

    presentation = Presentation(args.deck)
    required = parse_slide_spec(args.require_table_slides, len(presentation.slides))
    errors = []
    for number in required:
        slide = presentation.slides[number - 1]
        if not any(shape.has_table for shape in slide.shapes):
            errors.append(f"Slide {number} is missing a native table object")

    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(
        f"Structure valid: {len(presentation.slides)} slides; "
        f"native table required on {required or 'none'}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
