from __future__ import annotations

import importlib
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


REPO_DIR = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = REPO_DIR / "skill" / "make-ppt" / "assets" / "src-template"
VALIDATOR = (
    REPO_DIR / "skill" / "make-ppt" / "scripts" / "validate_pptx_structure.py"
)

sys.path.insert(0, str(TEMPLATE_DIR))
P = importlib.import_module("primitives")
T = importlib.import_module("theme")


def new_slide():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


class NativeTablePrimitiveTests(unittest.TestCase):
    def test_add_native_table_creates_one_real_table_object(self) -> None:
        _, slide = new_slide()

        frame = P.add_native_table(
            slide,
            Inches(0.6),
            Inches(2.2),
            Inches(12.1),
            Inches(3.8),
            headers=["技法", "落地"],
            rows=[
                ["冪等", "Idempotency-Key + notification_log"],
                ["併發控制", "If-Match / ETag + 樂觀鎖"],
            ],
            col_widths=[0.2, 0.8],
            row_label_col=0,
        )

        self.assertTrue(frame.has_table)
        self.assertEqual(1, sum(shape.has_table for shape in slide.shapes))
        self.assertEqual("技法", frame.table.cell(0, 0).text)
        self.assertEqual("If-Match / ETag + 樂觀鎖", frame.table.cell(2, 1).text)
        self.assertEqual(14, frame.table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size.pt)

    def test_structure_validator_rejects_fake_table_and_accepts_native_table(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fake_path = root / "fake.pptx"
            real_path = root / "real.pptx"

            fake_prs, fake_slide = new_slide()
            fake_slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5), Inches(4), Inches(0.5)
            ).text = "技法"
            fake_slide.shapes.add_textbox(
                Inches(4.5), Inches(0.5), Inches(4), Inches(0.5)
            ).text = "落地"
            fake_prs.save(fake_path)

            real_prs, real_slide = new_slide()
            P.add_native_table(
                real_slide,
                Inches(0.5),
                Inches(0.5),
                Inches(8),
                Inches(2),
                headers=["技法", "落地"],
                rows=[["冪等", "Idempotency-Key"]],
            )
            real_prs.save(real_path)

            fake = subprocess.run(
                [
                    sys.executable,
                    str(VALIDATOR),
                    str(fake_path),
                    "--require-table-slides",
                    "1",
                ],
                capture_output=True,
                text=True,
            )
            real = subprocess.run(
                [
                    sys.executable,
                    str(VALIDATOR),
                    str(real_path),
                    "--require-table-slides",
                    "1",
                ],
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(0, fake.returncode)
            self.assertIn("missing a native table object", fake.stderr)
            self.assertEqual(0, real.returncode, real.stderr)


class CoverPrimitiveTests(unittest.TestCase):
    @staticmethod
    def _shape_texts(slide) -> list[str]:
        return [
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text
        ]

    def test_editorial_light_cover_matches_reference_hierarchy(self) -> None:
        _, slide = new_slide()

        P.add_cover_slide(
            slide,
            [("從觀望\n到全公司落地", {})],
            variant="editorial-light",
            context="AI AGENTIC CODING · ADOPTION",
            eyebrow="導入實戰",
            subtitle="兩年 AI Agentic Coding 導入實戰\n流程整合與安全治理",
            presenter="appLeboy",
            date="2026.07.01",
            page_marker="01 / 24",
        )

        self.assertEqual(T.BG, slide.background.fill.fore_color.rgb)
        texts = self._shape_texts(slide)
        self.assertIn("從觀望\n到全公司落地", texts)
        self.assertIn("兩年 AI Agentic Coding 導入實戰\n流程整合與安全治理", texts)
        self.assertIn("appLeboy", texts)
        self.assertIn("2026.07.01", texts)
        title = next(
            shape
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and shape.text == "從觀望\n到全公司落地"
        )
        self.assertEqual(T.S_COVER_TITLE_LIGHT, title.text_frame.paragraphs[0].runs[0].font.size.pt)

    def test_report_dark_cover_has_no_fixed_footer_summary(self) -> None:
        _, slide = new_slide()

        P.add_cover_slide(
            slide,
            [("年度績效成果報告", {})],
            variant="report-dark",
            context="雲端處 ITSM 開發團隊",
            eyebrow="2026 · PERFORMANCE REVIEW",
            presenter="李相定",
            affiliation="架構師助手 / ITSM 開發",
            date="2026 · 07",
        )

        self.assertEqual(T.INK, slide.background.fill.fore_color.rgb)
        all_text = "\n".join(self._shape_texts(slide))
        self.assertIn("年度績效成果報告", all_text)
        self.assertNotIn("三件事", all_text)

    def test_report_dark_footer_highlight_is_truly_optional(self) -> None:
        _, without_footer = new_slide()
        P.add_cover_slide(
            without_footer,
            [("七月工作進度簡報", {})],
            variant="report-dark",
        )
        self.assertNotIn("回扣 KPI", "\n".join(self._shape_texts(without_footer)))

        _, with_footer = new_slide()
        P.add_cover_slide(
            with_footer,
            [("七月工作進度簡報", {})],
            variant="report-dark",
            footer_highlight="回扣 KPI 與跨月交付風險",
        )
        self.assertIn("回扣 KPI 與跨月交付風險", "\n".join(self._shape_texts(with_footer)))

    def test_unknown_cover_variant_is_rejected(self) -> None:
        _, slide = new_slide()
        with self.assertRaisesRegex(ValueError, "cover variant"):
            P.add_cover_slide(slide, [("封面", {})], variant="unknown")


if __name__ == "__main__":
    unittest.main()
