import sys
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches as I
from pptx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skill/make-ppt/assets/src-template'))
import relationship_layouts as R
import theme as T


class RelationshipTests(unittest.TestCase):
    def setUp(self):
        self.prs=Presentation()
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def test_matrix_preserves_states_and_native_table(self):
        frame=R.add_status_matrix(self.slide,I(1),I(1),I(10),I(3),
            headers=['能力','專案'],rows=[['A','yes'],['B','unknown']],
            states={'yes':dict(symbol='●',label='已確認',color=T.ORANGE),
                    'unknown':dict(symbol='?',label='未記載',color=T.MUTED)},
            groups=[dict(start=0,end=1,color=T.BLUE)])
        self.assertEqual(sum(s.has_table for s in self.slide.shapes),1)
        self.assertEqual(frame.table.cell(2,1).text,'?')
        # PowerPoint requires the line elements before fill in schema order.
        props=frame.table.cell(1,1)._tc.get_or_add_tcPr()
        self.assertEqual([node.tag for node in list(props)[:4]],
                         [qn(name) for name in ('a:lnL','a:lnR','a:lnT','a:lnB')])
        self.assertEqual(props.find(qn('a:lnL')).find(qn('a:solidFill'))[0].get('val'),
                         str(T.HAIRLINE))
        rail=self.slide.shapes[1]
        self.assertEqual(rail.top,frame.top+frame.table.rows[0].height)
        self.assertEqual(rail.height,sum(row.height for row in list(frame.table.rows)[1:]))

    def test_unknown_state_is_not_silently_treated_as_absence(self):
        with self.assertRaises(ValueError):
            R.add_status_matrix(self.slide,I(1),I(1),I(10),I(3),
                headers=['能力','專案'],rows=[['A',None]],
                states={'yes':dict(symbol='●',label='已確認',color=T.ORANGE)})

    def test_timeline_preserves_gap_and_proportional_spans(self):
        R.add_multitrack_timeline(self.slide,I(1),I(1),I(10),I(2),
            ticks=[(0,'0'),(10,'10')],tracks=[dict(label='A',color=T.BLUE,intervals=[(0,2),(6,10)])])
        bars=[s for s in self.slide.shapes if s.height==I(0.2)]
        self.assertEqual(len(bars),2)
        self.assertEqual(bars[1].width,2*bars[0].width)
        self.assertGreater(bars[1].left,bars[0].left+bars[0].width)

    def test_invalid_dates_and_dense_groups_fail(self):
        with self.assertRaises(ValueError):
            R.add_multitrack_timeline(self.slide,I(1),I(1),I(10),I(2),
                ticks=[(0,'0'),(10,'10')],tracks=[dict(label='A',color=T.BLUE,intervals=[(8,2)])])
        with self.assertRaises(ValueError):
            R.add_layered_list(self.slide,I(1),I(1),I(10),I(1),
                groups=[dict(label='A',color=T.BLUE,items=[('a','b')]*10)])

if __name__=='__main__':
    unittest.main()
