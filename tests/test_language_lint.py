from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


REPO_DIR = Path(__file__).resolve().parents[1]
LINTER = REPO_DIR / "skill" / "make-ppt" / "scripts" / "lint_zh_tw.py"


def run_linter(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(LINTER), *args],
        check=False,
        capture_output=True,
        text=True,
    )


class TaiwanLanguageLintTests(unittest.TestCase):
    def test_outline_mode_scans_visible_fields_but_not_source_quotes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            outline = Path(temporary) / "outline.md"
            outline.write_text(
                "# Presentation Specification\n\n"
                "## Slide 01 — 操作\n\n"
                "### Title\n\n雙擊項目後開啟配置頁面\n\n"
                "### Source Facts\n\nSupporting text: 雙擊是來源逐字引文\n",
                encoding="utf-8",
            )
            result = run_linter(str(outline), "--outline-visible-only")
            self.assertEqual(1, result.returncode)
            self.assertEqual(1, result.stdout.count("ERROR 雙擊"))
            self.assertIn("REVIEW 配置", result.stdout)

            allowed = run_linter(
                str(outline),
                "--outline-visible-only",
                "--allow-term",
                "雙擊",
            )
            self.assertEqual(0, allowed.returncode)
            self.assertIn("REVIEW 配置", allowed.stdout)

    def test_pptx_mode_finds_visible_shape_text(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            deck = Path(temporary) / "deck.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            box.text_frame.text = "播放視頻"
            presentation.save(deck)

            result = run_linter(str(deck))
            self.assertEqual(1, result.returncode)
            self.assertIn("ERROR 視頻", result.stdout)

    def test_clean_taiwan_wording_passes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            outline = Path(temporary) / "outline.md"
            outline.write_text("### Title\n\n按兩下項目後開啟設定頁面\n", encoding="utf-8")
            result = run_linter(str(outline), "--outline-visible-only")
            self.assertEqual(0, result.returncode)
            self.assertIn("check: pass", result.stdout)


if __name__ == "__main__":
    unittest.main()
